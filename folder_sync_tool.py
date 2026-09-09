"""
Folder Sync Comparison Tool
============================

Compares two folder trees -- which may live on two different machines
(e.g. one local, one reached over RDP or SFTP), and any mix of Windows and
Mac -- and produces an Excel report showing:

  1. Files/folders that exist in B but not A  -> copy B -> A
  2. Files/folders that exist in A but not B  -> copy A -> B
  3. Files that exist in both, where A's copy is newer -> copy A -> B
  4. Files that exist in both, where B's copy is newer -> copy B -> A

Subfolders are walked recursively on both sides.

This script works in two modes (set MODE in .env or below):

  "scan"            Walk ONE local folder and save its contents (relative
                     paths, sizes, last-modified times) to a small JSON
                     "manifest" file. Useful on its own as a point-in-time
                     inventory of a folder.

  "sftp_compare"    Folder A lives on an SFTP server, folder B is local to
                     the machine running this script. Connects to the SFTP
                     server, walks folder A remotely, walks folder B on
                     disk, and compares both in one run -- no manifest
                     files or copying needed.

Cross-platform (Windows <-> Mac) notes -- handled automatically below:
  - Both Windows (NTFS) and Mac (APFS/HFS+) treat filenames as case-
    insensitive by default, so "Report.pdf" and "report.pdf" are matched
    as the *same* file rather than flagged as missing on both sides --
    the exact-case difference is called out separately instead.
  - Mac has historically stored accented filenames (e.g. "café.txt") in
    Unicode NFD (decomposed) form, while Windows stores them NFC
    (composed). Same visible name, different bytes. Paths are Unicode-
    normalized before comparison so these aren't treated as different
    files.
  - Mac silently creates ".DS_Store" and "._filename" sidecar files when
    browsing or copying to a non-Mac drive. These (plus the Windows
    equivalents like Thumbs.db) are excluded automatically.

Configuration:
    All settings below are read from environment variables (loaded from a
    local .env file via python-dotenv), falling back to the defaults shown
    if a variable isn't set. Copy .env.example to .env and fill in your
    real paths/credentials -- .env is gitignored and never committed.

Requirements:
    pip install -r requirements.txt

Note on timestamps: comparisons use each file's raw epoch modified-time
(seconds since 1970, timezone-independent), so this works fine even if
the two machines are in different timezones -- as long as their clocks
are reasonably accurate. TIMESTAMP_TOLERANCE_SECONDS below absorbs small
clock drift / filesystem timestamp rounding differences.
"""

import json
import os
import stat
import unicodedata
from dataclasses import dataclass, asdict
from datetime import datetime
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

load_dotenv()


def _env_bool(name: str, default: bool) -> bool:
    return os.getenv(name, str(default)).strip().lower() in ("1", "true", "yes")


# ============================== CONFIG ==============================

MODE = os.getenv("MODE", "scan")   # "scan" | "sftp_compare"

# ---- used when MODE == "scan" ----
SCAN_SIDE_LABEL = os.getenv("SCAN_SIDE_LABEL", "A")     # "A" or "B" -- just a label for filenames
SCAN_ROOT_FOLDER = os.getenv("SCAN_ROOT_FOLDER", r"C:\path\to\folder_to_scan")
MANIFEST_OUTPUT_PATH = f"manifest_{SCAN_SIDE_LABEL}.json"

# ---- used when MODE == "sftp_compare" ----
# Folder A lives on the SFTP server; folder B is local to this machine.
SFTP_HOST = os.getenv("SFTP_HOST", "sftp.example.com")
SFTP_PORT = int(os.getenv("SFTP_PORT", "22"))
SFTP_USERNAME = os.getenv("SFTP_USERNAME", "")
SFTP_PASSWORD = os.getenv("SFTP_PASSWORD", "")              # leave unset if using a private key instead
SFTP_PRIVATE_KEY_PATH = os.getenv("SFTP_PRIVATE_KEY_PATH", "")   # e.g. C:\Users\you\.ssh\id_rsa
SFTP_ROOT_FOLDER = os.getenv("SFTP_ROOT_FOLDER", "/remote/path/to/folder_a")   # POSIX-style path on the server
LOCAL_ROOT_FOLDER = os.getenv("LOCAL_ROOT_FOLDER", r"C:\path\to\folder_b")     # path on THIS (client) machine

# ---- used by "sftp_compare" ----
LABEL_A = os.getenv("LABEL_A", "Folder A")
LABEL_B = os.getenv("LABEL_B", "Folder B")
EXCEL_REPORT_PATH = os.getenv("EXCEL_REPORT_PATH", "folder_sync_report.xlsx")

# Two files whose modified times differ by less than this are treated as
# "the same" rather than one being newer. Guards against clock skew /
# filesystem timestamp rounding between the two machines.
TIMESTAMP_TOLERANCE_SECONDS = int(os.getenv("TIMESTAMP_TOLERANCE_SECONDS", "2"))

# Match "Report.pdf" and "report.pdf" as the same file (recommended True,
# since both Windows/NTFS and Mac/APFS are case-insensitive by default).
# Case differences are still reported separately so you can rename if you
# care about exact casing.
CASE_INSENSITIVE_MATCH = _env_bool("CASE_INSENSITIVE_MATCH", True)

# Set True to also list files that are identical on both sides (no action
# needed) as their own sheet. Off by default to keep the report focused.
INCLUDE_IDENTICAL_SHEET = _env_bool("INCLUDE_IDENTICAL_SHEET", False)

# Names to ignore everywhere (junk/system files that shouldn't drive sync
# decisions). Matched by exact folder/file name.
EXCLUDE_NAMES = {
    # Windows
    "Thumbs.db", "desktop.ini", "$RECYCLE.BIN", "System Volume Information",
    # Mac
    ".DS_Store", ".AppleDouble", ".AppleDB", ".AppleDesktop",
    ".Spotlight-V100", ".Trashes", ".fseventsd", ".TemporaryItems",
    ".VolumeIcon.icns", ".com.apple.timemachine.donotpresent",
    # Version control / misc
    ".git", "__pycache__",
}

# Names *starting with* any of these are also ignored -- covers Mac's
# "._filename" AppleDouble sidecar files, created automatically whenever
# a Mac copies files to a non-Mac drive (network share, USB, etc.).
EXCLUDE_PREFIXES = ("._",)

# ======================================================================


@dataclass
class Entry:
    is_dir: bool
    size: int           # bytes; 0 for directories
    mtime_epoch: float  # seconds since epoch -- timezone independent
    mtime_human: str    # human-readable local time, for display only


def _is_excluded(name: str) -> bool:
    return name in EXCLUDE_NAMES or name.startswith(EXCLUDE_PREFIXES)


def normalize_path_for_matching(rel_path: str) -> str:
    """Key used to decide whether two paths refer to 'the same' file.

    Normalizes Unicode form (Mac historically stores accented names
    decomposed/NFD, Windows stores them composed/NFC) and, optionally,
    case, so equivalent filenames from different OSes line up.
    """
    normalized = unicodedata.normalize("NFC", rel_path)
    if CASE_INSENSITIVE_MATCH:
        normalized = normalized.lower()
    return normalized


def scan_folder(root: str) -> dict:
    """Walk `root` recursively and return {relative_posix_path: entry-dict}."""
    root_path = Path(root).resolve()
    if not root_path.is_dir():
        raise NotADirectoryError(f"Not a folder: {root}")

    manifest = {}
    for dirpath, dirnames, filenames in os.walk(root_path):
        # prune excluded dirs in place so os.walk doesn't descend into them
        dirnames[:] = [d for d in dirnames if not _is_excluded(d)]
        current = Path(dirpath)

        # record subfolders too, so empty / missing folders show up
        for d in dirnames:
            full = current / d
            rel = full.relative_to(root_path).as_posix()
            try:
                st = full.stat()
            except OSError:
                continue
            manifest[rel] = asdict(Entry(
                is_dir=True,
                size=0,
                mtime_epoch=st.st_mtime,
                mtime_human=datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
            ))

        for f in filenames:
            if _is_excluded(f):
                continue
            full = current / f
            rel = full.relative_to(root_path).as_posix()
            try:
                st = full.stat()
            except OSError:
                continue  # unreadable file (permissions, broken link, etc.)
            manifest[rel] = asdict(Entry(
                is_dir=False,
                size=st.st_size,
                mtime_epoch=st.st_mtime,
                mtime_human=datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
            ))

    return manifest


def connect_sftp(host: str, port: int, username: str, password: str, private_key_path: str):
    """Open an SFTP connection and return (sftp_client, transport).

    Caller is responsible for closing both (sftp_client.close(),
    transport.close()) when done.
    """
    import paramiko

    transport = paramiko.Transport((host, port))
    if private_key_path:
        pkey = paramiko.RSAKey.from_private_key_file(private_key_path)
        transport.connect(username=username, pkey=pkey)
    else:
        transport.connect(username=username, password=password)
    sftp = paramiko.SFTPClient.from_transport(transport)
    return sftp, transport


def scan_folder_sftp(sftp, root: str) -> dict:
    """Walk `root` recursively on the SFTP server and return the same
    {relative_posix_path: entry-dict} shape as scan_folder()."""
    root = root.rstrip("/") or "/"
    manifest = {}

    def _walk(remote_dir: str, rel_prefix: str) -> None:
        try:
            entries = sftp.listdir_attr(remote_dir)
        except IOError:
            return  # unreadable dir (permissions, broken symlink, etc.)

        for attr in entries:
            name = attr.filename
            if _is_excluded(name):
                continue
            remote_path = f"{remote_dir}/{name}"
            rel = f"{rel_prefix}/{name}" if rel_prefix else name
            is_dir = stat.S_ISDIR(attr.st_mode)
            manifest[rel] = asdict(Entry(
                is_dir=is_dir,
                size=0 if is_dir else attr.st_size,
                mtime_epoch=attr.st_mtime,
                mtime_human=datetime.fromtimestamp(attr.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
            ))
            if is_dir:
                _walk(remote_path, rel)

    _walk(root, "")
    return manifest


def save_manifest(manifest: dict, out_path: str, root: str) -> None:
    payload = {
        "scanned_root": str(Path(root).resolve()),
        "scanned_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "entries": manifest,
    }
    with open(out_path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, indent=2, ensure_ascii=False)
    print(f"Scanned {len(manifest)} item(s) from {root}")
    print(f"Manifest written to {Path(out_path).resolve()}")


def compare_manifests(manifest_a: dict, manifest_b: dict, tolerance: float) -> dict:
    # map normalized-key -> actual path, per side, so we can match
    # case/unicode-equivalent names while still displaying real names
    norm_a = {normalize_path_for_matching(p): p for p in manifest_a}
    norm_b = {normalize_path_for_matching(p): p for p in manifest_b}

    only_in_a, only_in_b = [], []
    a_newer, b_newer = [], []
    identical = []
    type_conflicts = []
    name_differences = []  # matched as the same file, but literal name differs

    for key in sorted(set(norm_a) | set(norm_b)):
        path_a = norm_a.get(key)
        path_b = norm_b.get(key)

        if path_a and not path_b:
            only_in_a.append((path_a, manifest_a[path_a]))
            continue
        if path_b and not path_a:
            only_in_b.append((path_b, manifest_b[path_b]))
            continue

        entry_a, entry_b = manifest_a[path_a], manifest_b[path_b]

        if path_a != path_b:
            name_differences.append((path_a, path_b))

        if entry_a["is_dir"] and entry_b["is_dir"]:
            continue  # folder exists both sides, nothing to copy itself

        if entry_a["is_dir"] != entry_b["is_dir"]:
            # same name is a file on one side, a folder on the other -- rare,
            # needs a human to look at it rather than an auto "copy" verdict
            type_conflicts.append((path_a, entry_a, entry_b))
            continue

        diff = entry_a["mtime_epoch"] - entry_b["mtime_epoch"]
        if diff > tolerance:
            a_newer.append((path_a, entry_a, entry_b))
        elif diff < -tolerance:
            b_newer.append((path_a, entry_a, entry_b))
        else:
            identical.append((path_a, entry_a, entry_b))

    return {
        "only_in_a": only_in_a, "only_in_b": only_in_b,
        "a_newer": a_newer, "b_newer": b_newer,
        "identical": identical, "type_conflicts": type_conflicts,
        "name_differences": name_differences,
    }


# ------------------------------ Excel report ------------------------------

HEADER_FILL = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
MISSING_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
NEWER_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
CONFLICT_FILL = PatternFill(start_color="D9D2E9", end_color="D9D2E9", fill_type="solid")
NOTE_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")


def _autosize(ws):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(length + 2, 12), 70)


def _header(ws, headers):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT


def _write_missing_sheet(wb, title, entries, source_label, fill):
    ws = wb.create_sheet(title)
    _header(ws, ["Relative Path", "Type", "Size (bytes)", f"Last Modified ({source_label})"])
    for rel_path, entry in entries:
        ws.append([
            rel_path,
            "Folder" if entry["is_dir"] else "File",
            "" if entry["is_dir"] else entry["size"],
            entry["mtime_human"],
        ])
        for cell in ws[ws.max_row]:
            cell.fill = fill
    ws.freeze_panes = "A2"
    _autosize(ws)


def _write_newer_sheet(wb, title, entries, newer_label, copy_to_label, fill):
    ws = wb.create_sheet(title)
    _header(ws, ["Relative Path", "Size A (bytes)", "Last Modified A",
                 "Size B (bytes)", "Last Modified B", "Action"])
    for rel_path, entry_a, entry_b in entries:
        ws.append([
            rel_path, entry_a["size"], entry_a["mtime_human"],
            entry_b["size"], entry_b["mtime_human"],
            f"Copy {newer_label} -> {copy_to_label}",
        ])
        for cell in ws[ws.max_row]:
            cell.fill = fill
    ws.freeze_panes = "A2"
    _autosize(ws)


def _write_type_conflict_sheet(wb, entries, fill):
    ws = wb.create_sheet("Type conflicts (file vs folder)")
    _header(ws, ["Relative Path", "Type in A", "Type in B"])
    for rel_path, entry_a, entry_b in entries:
        ws.append([rel_path, "Folder" if entry_a["is_dir"] else "File",
                   "Folder" if entry_b["is_dir"] else "File"])
        for cell in ws[ws.max_row]:
            cell.fill = fill
    ws.freeze_panes = "A2"
    _autosize(ws)


def _write_name_differences_sheet(wb, pairs, fill):
    ws = wb.create_sheet("Name differences (case/accents)")
    _header(ws, ["Name in A", "Name in B", "Note"])
    for path_a, path_b in pairs:
        ws.append([path_a, path_b, "Same file, different case or accent encoding"])
        for cell in ws[ws.max_row]:
            cell.fill = fill
    ws.freeze_panes = "A2"
    _autosize(ws)


def _write_identical_sheet(wb, entries):
    ws = wb.create_sheet("Identical (no action)")
    _header(ws, ["Relative Path", "Size (bytes)", "Last Modified"])
    for rel_path, entry_a, _ in entries:
        ws.append([rel_path, entry_a["size"], entry_a["mtime_human"]])
    ws.freeze_panes = "A2"
    _autosize(ws)


def write_excel_report(results, label_a, label_b, out_path):
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Summary", 0)
    ws.append(["Folder Sync Comparison Summary"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Folder A", label_a])
    ws.append(["Folder B", label_b])
    ws.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws.append([])
    _header(ws, ["Category", "Count"])
    ws.append(["Missing in A (copy from B)", len(results["only_in_b"])])
    ws.append(["Missing in B (copy from A)", len(results["only_in_a"])])
    ws.append(["A is newer (copy A -> B)", len(results["a_newer"])])
    ws.append(["B is newer (copy B -> A)", len(results["b_newer"])])
    ws.append(["Type conflicts (file vs folder)", len(results["type_conflicts"])])
    ws.append(["Name differences (case/accents)", len(results["name_differences"])])
    ws.append(["Identical on both sides", len(results["identical"])])
    _autosize(ws)

    _write_missing_sheet(wb, "Missing in A (copy from B)", results["only_in_b"], label_b, MISSING_FILL)
    _write_missing_sheet(wb, "Missing in B (copy from A)", results["only_in_a"], label_a, MISSING_FILL)
    _write_newer_sheet(wb, "A newer (copy A to B)", results["a_newer"], label_a, label_b, NEWER_FILL)
    _write_newer_sheet(wb, "B newer (copy B to A)", results["b_newer"], label_b, label_a, NEWER_FILL)

    if results["type_conflicts"]:
        _write_type_conflict_sheet(wb, results["type_conflicts"], CONFLICT_FILL)

    if results["name_differences"]:
        _write_name_differences_sheet(wb, results["name_differences"], NOTE_FILL)

    if INCLUDE_IDENTICAL_SHEET:
        _write_identical_sheet(wb, results["identical"])

    wb.save(out_path)
    print(f"Excel report written to {Path(out_path).resolve()}")


# ------------------------------ main ------------------------------

def main():
    if MODE == "scan":
        manifest = scan_folder(SCAN_ROOT_FOLDER)
        save_manifest(manifest, MANIFEST_OUTPUT_PATH, SCAN_ROOT_FOLDER)

    elif MODE == "sftp_compare":
        sftp, transport = connect_sftp(
            SFTP_HOST, SFTP_PORT, SFTP_USERNAME, SFTP_PASSWORD, SFTP_PRIVATE_KEY_PATH
        )
        try:
            manifest_a = scan_folder_sftp(sftp, SFTP_ROOT_FOLDER)
            print(f"Scanned {len(manifest_a)} item(s) from {SFTP_HOST}:{SFTP_ROOT_FOLDER}")
        finally:
            sftp.close()
            transport.close()

        manifest_b = scan_folder(LOCAL_ROOT_FOLDER)
        print(f"Scanned {len(manifest_b)} item(s) from {LOCAL_ROOT_FOLDER}")

        results = compare_manifests(manifest_a, manifest_b, TIMESTAMP_TOLERANCE_SECONDS)
        write_excel_report(results, LABEL_A, LABEL_B, EXCEL_REPORT_PATH)

    else:
        raise ValueError(f"Unknown MODE: {MODE!r}")


if __name__ == "__main__":
    main()
